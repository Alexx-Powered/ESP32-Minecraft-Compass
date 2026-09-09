from __future__ import annotations

import argparse
import re
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# compact_reference_guide preset (exact controlled tokens)
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
BASE_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"  # Named glyph-coverage override for Chinese text.
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "17231F"
MUTED = "66746E"
GREEN = "123D32"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
GRID = "B9C5BE"


def set_run_font(run, name: str = BASE_FONT, size: float | None = None,
                 color: str | None = None, bold: bool | None = None,
                 italic: bool | None = None, east_asia: str = CJK_FONT) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS_DXA.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), GRID)


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths sum to {sum(widths)}, expected {CONTENT_WIDTH_DXA}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[index] / 1440)
            set_cell_margins(cell)

    set_table_borders(table)


def paragraph_bottom_border(paragraph, color: str = GREEN, size: int = 16,
                            space: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BASE_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), BASE_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BASE_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in tokens.items():
        style = doc.styles[name]
        style.font.name = BASE_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), BASE_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BASE_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    for style_name in ("Header", "Footer"):
        style = doc.styles[style_name]
        style.font.name = BASE_FONT
        style.font.size = Pt(8.5)
        style.font.color.rgb = RGBColor.from_string(MUTED)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    code = doc.styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(8.3)
    code.font.color.rgb = RGBColor.from_string(INK)
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    code.paragraph_format.left_indent = Inches(0.12)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.05

    caption = doc.styles["Caption"]
    caption.font.name = BASE_FONT
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)


def setup_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), 2)
    left = p.add_run("MinecraftCompass · 技术方案")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    right = p.add_run("\tTarget / Geofence · v1.0")
    set_run_font(right, size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    run = fp.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    page_run = fp.add_run()
    set_run_font(page_run, size=8.5, color=MUTED)
    add_field(page_run, "PAGE")
    run = fp.add_run(" 页 / 共 ")
    set_run_font(run, size=8.5, color=MUTED)
    total_run = fp.add_run()
    set_run_font(total_run, size=8.5, color=MUTED)
    add_field(total_run, "NUMPAGES")
    run = fp.add_run(" 页")
    set_run_font(run, size=8.5, color=MUTED)


def add_masthead(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(16)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("TARGET 与电子围栏配置方案")
    set_run_font(run, size=23, color="000000", bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.keep_with_next = True
    run = subtitle.add_run("MinecraftCompass · ESP32-S3 本地配置、NVS、坐标系与高德生成器设计")
    set_run_font(run, size=14, color="373737")

    metadata = [
        ("面向", "MinecraftCompass 项目开发与审阅"),
        ("版本", "方案 v1.0 / JSON schema 1"),
        ("日期", "2026-08-28"),
        ("范围", "一个 Target、最多两个电子围栏、WPA2 SoftAP、本地页面与 NVS"),
        ("状态", "代码已实现；ESP-IDF 5.5.5 编译与镜像大小检查通过"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        lr = p.add_run(f"{label}：")
        set_run_font(lr, size=11, color="000000", bold=True)
        vr = p.add_run(value)
        set_run_font(vr, size=11, color="000000")

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(12)
    paragraph_bottom_border(rule)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), BASE_FONT)
    r_fonts.set(qn("w:hAnsi"), BASE_FONT)
    r_fonts.set(qn("w:eastAsia"), CJK_FONT)
    r_pr.extend([r_fonts, color, underline])
    run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)


def add_inline(paragraph, text: str, base_size: float = 11,
               color: str = INK) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=base_size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, color=color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=base_size - 0.7,
                         color=DARK_BLUE, east_asia=CJK_FONT)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF2F4")
            run._r.get_or_add_rPr().append(shd)
        else:
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=base_size, color=color)


def create_numbering(doc: Document, bullet: bool,
                     start_value: int = 1) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId")))
                    for node in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), str(start_value))
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "271")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, indent, spacing])
    level.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
    if bullet:
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Symbol")
        fonts.set(qn("w:hAnsi"), "Symbol")
        r_pr.append(fonts)
        level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)

    nums = [int(node.get(qn("w:numId")))
            for node in numbering.findall(qn("w:num"))]
    num_id = max(nums, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc: Document, text: str, num_id: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)
    add_inline(p, text)


def column_widths(headers: list[str]) -> list[int]:
    count = len(headers)
    first = headers[0] if headers else ""
    if count == 2:
        return [2700, 6660]
    if count == 3 and first == "字段":
        return [1550, 2800, 5010]
    if count == 3 and first == "命名空间":
        return [1850, 1600, 5910]
    if count == 3 and first == "场景":
        return [1900, 3000, 4460]
    if count == 5:
        return [1500, 1050, 650, 2860, 3300]
    base = CONTENT_WIDTH_DXA // count
    result = [base] * count
    result[-1] += CONTENT_WIDTH_DXA - sum(result)
    return result


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    widths = column_widths(rows[0])
    set_table_geometry(table, widths)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, HEADER_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            add_inline(p, value, base_size=9.3,
                       color=INK)
            if row_index == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        if row_index == 0:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        # Keep a logical row intact across pages. This avoids blank-looking row
        # continuations and makes repeated headers easier to scan in LibreOffice.
        tr_pr = table.rows[row_index]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        tr_pr.append(cant_split)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def shade_paragraph(paragraph, fill: str = LIGHT_FILL) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "4")
        node.set(qn("w:color"), "D8E0DB")
        p_bdr.append(node)
    p_pr.append(p_bdr)


def wrap_code(code: str, width: int = 88) -> str:
    output: list[str] = []
    for line in code.splitlines():
        if len(line) <= width:
            output.append(line)
            continue
        indent = len(line) - len(line.lstrip())
        wrapped = textwrap.wrap(
            line,
            width=width,
            subsequent_indent=" " * min(indent + 2, 16),
            replace_whitespace=False,
            drop_whitespace=False,
            break_long_words=True,
            break_on_hyphens=False,
        )
        output.extend(wrapped or [line])
    return "\n".join(output)


def add_code_block(doc: Document, code: str) -> None:
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    shade_paragraph(p)
    run = p.add_run(wrap_code(code))
    set_run_font(run, name="Consolas", size=8.3, color=INK,
                 east_asia=CJK_FONT)


def set_image_alt_text(inline_shape, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def add_image(doc: Document, image_path: Path, alt: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = True
    run = p.add_run()
    width = 5.95 if "任务优先级" in alt else 6.35
    shape = run.add_picture(str(image_path), width=Inches(width))
    set_image_alt_text(shape, alt)
    caption = doc.add_paragraph(style="Caption")
    caption.add_run(f"图：{alt}")


def add_callout(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.14)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    shade_paragraph(p, "EAF4EE")
    add_inline(p, text, color=GREEN)


def parse_markdown(doc: Document, markdown_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0
    skipped_title = False
    list_type: str | None = None
    list_num_id: int | None = None

    def reset_list() -> None:
        nonlocal list_type, list_num_id
        list_type = None
        list_num_id = None

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            reset_list()
            index += 1
            continue
        if stripped.startswith("# ") and not skipped_title:
            skipped_title = True
            index += 1
            continue
        if stripped.startswith("> 文档版本"):
            while index < len(lines) and lines[index].lstrip().startswith(">"):
                index += 1
            continue
        if stripped.startswith("```"):
            reset_list()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            index += 1
            add_code_block(doc, "\n".join(code_lines))
            continue
        image_match = re.fullmatch(r"!\[([^\]]+)\]\(([^)]+)\)", stripped)
        if image_match:
            reset_list()
            image_path = (markdown_path.parent / image_match.group(2)).resolve()
            add_image(doc, image_path, image_match.group(1))
            index += 1
            continue
        if stripped.startswith("|"):
            reset_list()
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows: list[list[str]] = []
            for row_index, table_line in enumerate(table_lines):
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                if row_index == 1 and all(re.fullmatch(r":?-{3,}:?", cell)
                                          for cell in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            reset_list()
            level = len(heading.group(1)) - 1
            p = doc.add_paragraph(style=f"Heading {level}")
            if level == 1 and (
                heading.group(2).startswith("4. JSON") or
                heading.group(2).startswith("9. FreeRTOS") or
                heading.group(2).startswith("12. 已知")
            ):
                p.paragraph_format.page_break_before = True
            add_inline(p, heading.group(2),
                       base_size={1: 16, 2: 13, 3: 12}[level],
                       color=BLUE if level < 3 else DARK_BLUE)
            index += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        number = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet or number:
            kind = "bullet" if bullet else "number"
            item_text = (bullet or number).group(1)
            if list_type != kind:
                list_type = kind
                start_value = 1 if bullet else int(re.match(r"^(\d+)", stripped).group(1))
                list_num_id = create_numbering(
                    doc, bullet=kind == "bullet", start_value=start_value)
            add_list_item(doc, item_text, list_num_id)
            index += 1
            continue
        if stripped == "---":
            reset_list()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            paragraph_bottom_border(p, color="C8D1CC", size=6, space=2)
            index += 1
            continue
        if stripped.startswith(">"):
            reset_list()
            add_callout(doc, stripped.lstrip("> "))
            index += 1
            continue

        reset_list()
        p = doc.add_paragraph()
        if stripped == "相关高德官方资料：":
            p.paragraph_format.page_break_before = True
        if stripped.endswith("："):
            p.paragraph_format.keep_with_next = True
        add_inline(p, stripped)
        index += 1


def set_document_language(doc: Document) -> None:
    settings = doc.settings._element
    theme_font_lang = settings.find(qn("w:themeFontLang"))
    if theme_font_lang is None:
        theme_font_lang = OxmlElement("w:themeFontLang")
        settings.append(theme_font_lang)
    theme_font_lang.set(qn("w:val"), "zh-CN")
    theme_font_lang.set(qn("w:eastAsia"), "zh-CN")


def build(markdown_path: Path, output_path: Path) -> None:
    doc = Document()
    setup_page(doc)
    setup_styles(doc)
    set_document_language(doc)
    doc.core_properties.title = "MinecraftCompass Target 与电子围栏配置方案"
    doc.core_properties.subject = "ESP32-S3 本地配置、NVS、坐标系、高德 JS API 与 FreeRTOS 评估"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "ESP-IDF, ESP32-S3, Target, Geofence, NVS, AMap, FreeRTOS"
    add_masthead(doc)
    parse_markdown(doc, markdown_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("markdown", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.markdown.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
